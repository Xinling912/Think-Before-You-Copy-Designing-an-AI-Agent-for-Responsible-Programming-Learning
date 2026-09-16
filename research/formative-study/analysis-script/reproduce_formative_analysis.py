"""Reproduce the formative-study quantitative tables and paper figures.

The public 40-participant CSV is the canonical analysis input. Part B ratings
are stored as numeric values; Part C1-C4 cells contain both a rating and a
written explanation. This script extracts the ratings, validates the complete
case dataset, and reproduces the quantitative evidence reported in the paper.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


SCRIPT_DIR = Path(__file__).resolve().parent
FORMATIVE_DIR = SCRIPT_DIR.parent
DEFAULT_INPUT = FORMATIVE_DIR / "Original Data" / "formative_responses_complete_40.csv"
DEFAULT_OUTPUT = SCRIPT_DIR / "reproduced_output"

B_LABELS = {
    "B1": "Turning to AI unconsciously",
    "B2": "Asking AI before trying",
    "B3": "Switching to another AI tool",
    "B4": "Relying on AI for debugging",
    "B5": "Concerned about AI use",
    "B6": "Hard to work without AI",
    "B7": "Less confident without AI",
    "B8": "Reduced problem-solving ability",
}

C_LABELS = {
    "C1.1": "Efficiency",
    "C1.2": "Improved performance",
    "C1.3": "Effective learning",
    "C1.4": "Easier learning",
    "C1.5": "Overall benefit",
    "C2.1": "Undermines education value",
    "C2.2": "Limits social interaction",
    "C2.3": "Hinders skill development",
    "C2.4": "Over-reliance",
    "C3.1": "Helpfulness",
    "C3.2": "Continued use",
    "C3.3": "Frequent use",
    "C3.4": "Recommendation",
    "C4.1": "Correct code",
    "C4.2": "Answer questions",
    "C4.3": "Code examples",
    "C4.4": "Advice/resources",
    "C4.5": "Explain concepts",
}

B_DIMENSIONS = [
    (
        "B1-B4",
        "Dependency behaviors",
        "Behavioral dependency patterns when using generative AI for programming.",
        ["B1", "B2", "B3", "B4"],
    ),
    (
        "B5-B8",
        "Negative consequences",
        "Concerns, reduced confidence, difficulty without AI, and perceived problem-solving impact.",
        ["B5", "B6", "B7", "B8"],
    ),
]

C_DIMENSIONS = [
    (
        "C1",
        "Perceived usefulness",
        "Perceived usefulness of AI tools for programming learning.",
        ["C1.1", "C1.2", "C1.3", "C1.4", "C1.5"],
    ),
    (
        "C2",
        "Responsible risks",
        "Perceived educational, social, transferable-skill, and over-reliance risks.",
        ["C2.1", "C2.2", "C2.3", "C2.4"],
    ),
    ("C3.1", "Helpfulness", "Perceived helpfulness for programming learning.", ["C3.1"]),
    ("C3.2", "Continued use", "Intention to continue using AI tools.", ["C3.2"]),
    ("C3.3", "Frequent use", "Intention to use AI tools frequently.", ["C3.3"]),
    ("C3.4", "Recommendation", "Willingness to recommend AI tools to peers.", ["C3.4"]),
    ("C4.1", "Correct code", "Expected support for correcting code.", ["C4.1"]),
    ("C4.2", "Answer questions", "Expected support for answering programming questions.", ["C4.2"]),
    ("C4.3", "Code examples", "Expected support through code examples.", ["C4.3"]),
    ("C4.4", "Advice/resources", "Expected learning advice and resources.", ["C4.4"]),
    ("C4.5", "Explain concepts", "Expected explanation of programming concepts.", ["C4.5"]),
]


def question_id(column: str) -> str | None:
    match = re.match(r"^(B\d+|C\d+\.\d+)\.", column)
    return match.group(1) if match else None


def question_text(column: str) -> str:
    return re.sub(r"^(?:B\d+|C\d+\.\d+)\.\s*", "", column).strip()


def parse_rating(value) -> int | None:
    text = str(value).strip()
    if re.fullmatch(r"[1-5](?:\.0)?", text):
        return int(float(text))
    match = re.search(r"Rating:\s*([1-5])", text, flags=re.IGNORECASE)
    return int(match.group(1)) if match else None


def item_columns(
    frame: pd.DataFrame,
    expected_items: list[str],
) -> dict[str, str]:
    available: dict[str, str] = {}
    for column in frame.columns:
        item = question_id(column)
        if item:
            available[item] = column
    missing = [item for item in expected_items if item not in available]
    if missing:
        raise ValueError(f"Missing expected Likert columns: {', '.join(missing)}")
    return {item: available[item] for item in expected_items}


def summarize_items(
    frame: pd.DataFrame,
    columns: dict[str, str],
    labels: dict[str, str],
    scale: str,
) -> pd.DataFrame:
    rows = []
    for item, column in columns.items():
        scores = frame[column].map(parse_rating).dropna().astype(int)
        counts = scores.value_counts().reindex(range(1, 6), fill_value=0)
        n = int(scores.size)
        if n != len(frame):
            raise ValueError(
                f"{item} has {n} valid ratings; expected {len(frame)} complete ratings"
            )
        row = {
            "Scale": scale,
            "Item": item,
            "Theme_Label": labels[item],
            "Question": question_text(column),
            "N": n,
            "Mean": round(float(scores.mean()), 3),
            "SD": round(float(scores.std(ddof=1)), 3),
            "Negative_pct_1_2": round(float((scores <= 2).mean()), 5),
            "Neutral_pct_3": round(float((scores == 3).mean()), 5),
            "Positive_pct_4_5": round(float((scores >= 4).mean()), 5),
        }
        for rating in range(1, 6):
            row[f"Rating_{rating}_count"] = int(counts[rating])
            row[f"Rating_{rating}_pct"] = round(float(counts[rating] / n), 5)
        rows.append(row)
    return pd.DataFrame(rows)


def aggregate_dimensions(
    item_table: pd.DataFrame,
    groups: list[tuple[str, str, str, list[str]]],
    scale: str,
) -> pd.DataFrame:
    source = item_table.set_index("Item")
    rows = []
    for item, label, description, members in groups:
        counts = {
            rating: int(
                sum(source.loc[member, f"Rating_{rating}_count"] for member in members)
            )
            for rating in range(1, 6)
        }
        values = [
            rating
            for rating in range(1, 6)
            for _ in range(counts[rating])
        ]
        scores = pd.Series(values, dtype="float64")
        n = int(scores.size)
        rows.append(
            {
                "Scale": scale,
                "Item": item,
                "Theme Label": label,
                "Question": description,
                "N": n,
                "Mean": round(float(scores.mean()), 3),
                "SD": round(float(scores.std(ddof=1)), 3),
                **{str(rating): counts[rating] for rating in range(1, 6)},
                "Negative % (1-2)": round((counts[1] + counts[2]) / n, 5),
                "Neutral % (3)": round(counts[3] / n, 5),
                "Positive % (4-5)": round((counts[4] + counts[5]) / n, 5),
            }
        )
    return pd.DataFrame(rows)


def item_table_for_workbook(table: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, source in table.iterrows():
        row = {
            "Scale": source["Scale"],
            "Item": source["Item"],
            "Theme Label": source["Theme_Label"],
            "Question": source["Question"],
            "N": source["N"],
            "Mean": source["Mean"],
            "SD": source["SD"],
        }
        for rating in range(1, 6):
            row[str(rating)] = source[f"Rating_{rating}_count"]
            row[f"{rating} %"] = source[f"Rating_{rating}_pct"]
        row["Negative % (1-2)"] = source["Negative_pct_1_2"]
        row["Neutral % (3)"] = source["Neutral_pct_3"]
        row["Positive % (4-5)"] = source["Positive_pct_4_5"]
        rows.append(row)
    return pd.DataFrame(rows)


def dimension_table_for_workbook(table: pd.DataFrame) -> pd.DataFrame:
    result = table.copy()
    for rating in range(1, 6):
        result.insert(
            result.columns.get_loc(str(rating)) + 1,
            f"{rating} %",
            (result[str(rating)] / result["N"]).round(5),
        )
    return result


def write_workbook(
    output: Path,
    b_items: pd.DataFrame,
    c_items: pd.DataFrame,
    b_dimensions: pd.DataFrame,
    c_dimensions: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        item_table_for_workbook(b_items).to_excel(
            writer, sheet_name="B_Quantitative", index=False
        )
        dimension_table_for_workbook(b_dimensions).to_excel(
            writer, sheet_name="B_Dimension_Quantitative", index=False
        )
        item_table_for_workbook(c_items).to_excel(
            writer, sheet_name="C_Quantitative", index=False
        )
        dimension_table_for_workbook(c_dimensions).to_excel(
            writer, sheet_name="C_Dimension_Quantitative", index=False
        )

    workbook = load_workbook(output)
    header_fill = PatternFill("solid", fgColor="E8E2FF")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="4F35C8")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            values = [str(cell.value or "") for cell in column_cells[:100]]
            width = min(max(max(map(len, values), default=8) + 2, 10), 48)
            sheet.column_dimensions[letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(output)


def register_font() -> str:
    candidates = [
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            pdfmetrics.registerFont(TTFont("OSFFont", str(path)))
            return "OSFFont"
    return "Helvetica"


def legend_layout(width: float, font: str, font_size: float) -> list[tuple[int, str, float]]:
    items = [
        (1, "1 = Strongly disagree"),
        (2, "2"),
        (3, "3"),
        (4, "4"),
        (5, "5 = Strongly agree"),
    ]
    left = 0.35 * cm
    right = width - 0.35 * cm
    return [
        (score, label, left + index * (right - left) / 4)
        for index, (score, label) in enumerate(items)
    ]


def draw_legend(
    pdf: canvas.Canvas,
    font: str,
    width: float,
    y: float,
    palette: dict[int, colors.Color],
) -> None:
    pdf.setFont(font, 6.5)
    for score, label, x in legend_layout(width, font, 6.5):
        pdf.setFillColor(palette[score])
        pdf.circle(x, y + 0.06 * cm, 0.065 * cm, fill=1, stroke=0)
        pdf.setFillColor(colors.black)
        if score == 5:
            pdf.drawRightString(width - 0.35 * cm, y - 0.01 * cm, label)
        else:
            pdf.drawString(x + 0.12 * cm, y - 0.01 * cm, label)


def draw_likert_pdf(table: pd.DataFrame, output: Path, width_cm: float, height_cm: float) -> None:
    font = register_font()
    width, height = width_cm * cm, height_cm * cm
    pdf = canvas.Canvas(str(output), pagesize=(width, height))
    palette = {
        1: colors.HexColor("#E8708D"),
        2: colors.HexColor("#F2A5B6"),
        3: colors.HexColor("#D9D9D9"),
        4: colors.HexColor("#96A8E8"),
        5: colors.HexColor("#3F55B7"),
    }
    margin_left = 3.65 * cm
    margin_right = 0.45 * cm
    center_x = margin_left + (width - margin_left - margin_right) * 0.39
    max_half = (width - margin_left - margin_right) * 0.48
    row_height = 0.48 * cm
    bar_height = 0.24 * cm
    first_y = height - 0.52 * cm
    last_y = first_y - (len(table) - 1) * row_height
    legend_y = max(0.18 * cm, last_y - 0.58 * cm)

    pdf.setStrokeColor(colors.HexColor("#999999"))
    pdf.setDash(2, 2)
    pdf.line(center_x, last_y - 0.03 * cm, center_x, first_y + 0.18 * cm)
    pdf.setDash()

    for index, (_, row) in enumerate(table.iterrows()):
        y = first_y - index * row_height
        n = int(row["N"])
        counts = {rating: int(row[str(rating)]) for rating in range(1, 6)}
        pdf.setFillColor(colors.black)
        pdf.setFont(font, 6.7)
        pdf.drawString(0.35 * cm, y + 0.03 * cm, str(row["Theme Label"]))

        left_cursor = center_x
        for rating in [3, 2, 1]:
            fraction = counts[rating] / n / 2 if rating == 3 else counts[rating] / n
            bar_width = fraction * max_half
            left_cursor -= bar_width
            pdf.setFillColor(palette[rating])
            pdf.rect(left_cursor, y, bar_width, bar_height, fill=1, stroke=0)

        right_cursor = center_x
        for rating in [3, 4, 5]:
            fraction = counts[rating] / n / 2 if rating == 3 else counts[rating] / n
            bar_width = fraction * max_half
            pdf.setFillColor(palette[rating])
            pdf.rect(right_cursor, y, bar_width, bar_height, fill=1, stroke=0)
            right_cursor += bar_width

        negative = (counts[1] + counts[2]) / n * 100
        positive = (counts[4] + counts[5]) / n * 100
        pdf.setFillColor(colors.black)
        pdf.setFont(font, 6.3)
        pdf.drawRightString(left_cursor - 0.05 * cm, y + 0.03 * cm, f"{negative:.1f}%")
        pdf.drawString(right_cursor + 0.05 * cm, y + 0.03 * cm, f"{positive:.1f}%")

    draw_legend(pdf, font, width, legend_y, palette)
    pdf.save()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    frame = pd.read_csv(args.input, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    if len(frame) != 40:
        raise ValueError(f"Expected 40 complete cases, found {len(frame)}")

    b_items = summarize_items(
        frame, item_columns(frame, list(B_LABELS)), B_LABELS, "B"
    )
    c_items = summarize_items(
        frame, item_columns(frame, list(C_LABELS)), C_LABELS, "C"
    )
    b_dimensions = aggregate_dimensions(b_items, B_DIMENSIONS, "B")
    c_dimensions = aggregate_dimensions(c_items, C_DIMENSIONS, "C")

    pd.concat([b_items, c_items], ignore_index=True).to_csv(
        args.output_dir / "BC_rating_frequency_counts_combined.csv",
        index=False,
        encoding="utf-8-sig",
    )
    pd.concat([b_dimensions, c_dimensions], ignore_index=True).to_csv(
        args.output_dir / "BC_dimension_rating_counts_combined.csv",
        index=False,
        encoding="utf-8-sig",
    )
    write_workbook(
        args.output_dir / "formative_analysis_tables.xlsx",
        b_items,
        c_items,
        b_dimensions,
        c_dimensions,
    )
    draw_likert_pdf(
        b_dimensions,
        args.output_dir / "B_likert_dimension_results.pdf",
        width_cm=11.4,
        height_cm=2.15,
    )
    draw_likert_pdf(
        c_dimensions,
        args.output_dir / "C_likert_dimension_results.pdf",
        width_cm=11.8,
        height_cm=6.25,
    )
    print(f"Reproduced analysis written to {args.output_dir}")


if __name__ == "__main__":
    main()
